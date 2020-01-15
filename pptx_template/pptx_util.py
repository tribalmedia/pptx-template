# utilities to access python-pptx's internal structure.
# coding=utf-8

from pptx.chart.axis import ValueAxis

import logging

log = logging.getLogger()

def set_value_axis(chart, max = None, min = None, is_second_axis = False):
    """
     2軸チャートの軸の最大値、最小値を設定する
    """
    axis = ValueAxis(chart._chartSpace.valAx_lst[1 if is_second_axis else 0])

    if max:
        axis.maximum_scale = float(max)
    if min:
        axis.minimum_scale = float(min)


def remove_slide(presentation, slide):
    """
     presentation から 指定した slide を削除する
    """
    id = [ (i, s.rId) for i,s in enumerate(presentation.slides._sldIdLst) if s.id == slide.slide_id ][0]
    log.debug(u"removing slide #%d %s (rel_id: %s)" % (id[0], slide.slide_id, id[1]))
    presentation.part.drop_rel(id[1])
    del presentation.slides._sldIdLst[id[0]]


def add_styled_run(paragraph, styled_run):
    """
     pptx.text.text._Paragraph にスタイルされた pptx.text.text._Run を追加する
    """
    run = paragraph.add_run()
    font = run.font
    font.bold = styled_run.font.bold
    font.color.rgb = styled_run.font.color.rgb
    font.italic = styled_run.font.italic
    font.language_id = styled_run.font.language_id
    font.name = styled_run.font.name
    font.size = styled_run.font.size
    font.underline = styled_run.font.underline
    return run
